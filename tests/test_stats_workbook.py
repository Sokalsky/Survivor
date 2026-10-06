import csv
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from survivor.stats_workbook import TableParser, add_sheet, download, parse_history, parse_projections, write_csv
from survivor.catalog import STAT_HEADERS, load_aliases, resolve_name
from survivor.import_stats import validate_csv


def historical_row(name='Test Player', player_id='testpl01', team='ABC', games=10):
    values={'name_display':name,'age':'25','team_name_abbr':team,'pos':'C','games':str(games),
            'mp':'300','pts':'150','trb':'50','ast':'20','stl':'10','blk':'5','fg3':'10',
            'fg':'60','fga':'120','ft':'20','fta':'25'}
    return '<tr>'+''.join(f'<td data-stat="{key}" '+(f'data-append-csv="{player_id}"' if key=='name_display' else '')+f'>{value}</td>' for key,value in values.items())+'</tr>'


def history_html(body):
    return '<title>2025-26 NBA Player Stats: Totals</title><table id="totals_stats">'+body+'</table>'


def projection_html():
    labels=['Player','PTS','REB','AST','BLK','STL','FG%','FT%','3PM','GP','MIN','TO']
    values=['1,500','500','200','50','100','.500','.800','100','50','1,500','100']
    return ('<title>NBA Fantasy Basketball Overall 2026-27 Projections</title>'
            '<time datetime="2026-09-19 12:00:00"></time><table id="data"><tr>'+
            ''.join('<th>'+label+'</th>' for label in labels)+'</tr><tr><td>'+
            '<a href="/nba/players/test-player.php" fp-player-name="Test Player">Test Player</a>'+
            ' <small>(TEST - PF,C)</small></td>'+''.join('<td>'+v+'</td>' for v in values)+'</tr></table>')


class StatsWorkbookTests(unittest.TestCase):
    def test_provider_suffixes_match_existing_archive_identities(self):
        aliases=load_aliases()
        self.assertEqual(resolve_name('P.J. Washington Jr.',aliases)[0],'pjwashington')
        self.assertEqual(resolve_name('Bobby Portis Jr.',aliases)[0],'bobbyportis')

    def test_aggregate_rows_replace_team_splits_and_postseason_is_excluded(self):
        html=history_html(historical_row(team='2TM')+historical_row(team='ABC',games=6)+historical_row(team='DEF',games=4))
        html+='<table id="totals_stats_post">'+historical_row(name='Postseason Only',player_id='post01')+'</table>'
        rows=parse_history(html,'2025-26')
        self.assertEqual(len(rows),1)
        self.assertEqual(rows[0]['games'],10)
        self.assertEqual(rows[0]['source_rows'],3)
        self.assertEqual((rows[0]['pts_pg'],rows[0]['fgm_pg'],rows[0]['ft_pct']),(15,6,.8))
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'actuals.csv'
            write_csv(path,rows,STAT_HEADERS)
            self.assertEqual(len(validate_csv(path)),1)

    def test_old_tot_label_and_comment_wrapped_table(self):
        table=history_html(historical_row(team='TOT')+historical_row(team='ABC'))
        html=table.replace('<table','<!-- <table').replace('</table>','</table> -->')
        self.assertEqual(len(parse_history(html,'2025-26')),1)

    def test_unknown_layout_wrong_season_and_missing_totals_fail(self):
        with self.assertRaises(ValueError):
            parse_history(history_html(historical_row()),'2024-25')
        with self.assertRaises(ValueError):
            parse_history(history_html(historical_row()+historical_row(team='DEF')),'2025-26')
        with self.assertRaises(ValueError):
            parse_history(history_html(historical_row().replace('>120<','><')),'2025-26')

    def test_projections_convert_totals_and_preserve_missing_volume(self):
        records,date=parse_projections(projection_html(),'2026-27')
        self.assertEqual(date,'2026-09-19')
        row=records[0]
        self.assertEqual((row['games'],row['pts_pg'],row['minutes_pg']),(50,30,30))
        self.assertEqual((row['nba_team'],row['positions'],row['fg_pct']),('TEST','PF,C',.5))
        self.assertFalse(row['ready_for_valuation'])
        self.assertTrue(all(row[field] is None for field in ('fgm_pg','fga_pg','ftm_pg','fta_pg')))
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'projection.csv'
            write_csv(path,records,STAT_HEADERS)
            with self.assertRaisesRegex(ValueError,'fgm_pg must be numeric'):
                validate_csv(path)
        with self.assertRaises(ValueError):
            parse_projections(projection_html(),'2025-26')
        with self.assertRaises(ValueError):
            parse_projections(projection_html().replace('<th>PTS</th>','<th>POINTS</th>'),'2026-27')

    def test_excel_preserves_blank_cells_numeric_types_and_literal_text(self):
        wb=Workbook();wb.remove(wb.active)
        add_sheet(wb,'Projections',[{'player':'=unsafe()','fga_pg':None,'fg_pct':.5}],['player','fga_pg','fg_pct'])
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'test.xlsx';wb.save(path)
            loaded=load_workbook(path)
            sheet=loaded['Projections']
            self.assertEqual(sheet['A2'].data_type,'s')
            self.assertIsNone(sheet['B2'].value)
            self.assertEqual(sheet['C2'].number_format,'0.0%')
            self.assertEqual(len(sheet.tables),1)
            loaded.close()

    def test_cached_source_hash_is_checked_without_network_access(self):
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder)
            (root/'test.html').write_text('modified',encoding='utf-8')
            (root/'test.json').write_text(json.dumps({'url':'https://example.test','sha256':'wrong'}),encoding='utf-8')
            with patch('survivor.stats_workbook.urlopen') as request:
                with self.assertRaisesRegex(ValueError,'provenance'):
                    download('https://example.test',root,'test')
                request.assert_not_called()


if __name__=='__main__':
    unittest.main()

"""Reproducibly catalog the source workbook, retaining original cells and notes."""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
import csv
from datetime import datetime, timezone
from decimal import Decimal
import hashlib
import json
from pathlib import Path
import re
import unicodedata

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from survivor.database import preview_database, railway_database

ROOT = Path(__file__).resolve().parents[1]
HEADER = re.compile(r'^year\s+(\d+|one|two)\s*-\s*(\d+)', re.I)
COLUMN_HEADERS = {'drafted player', 'final roster', 'cost', 'year'}


def now():
    return datetime.now(timezone.utc).isoformat()


def name_key(name):
    text = unicodedata.normalize('NFKD', str(name)).casefold()
    return ''.join(c for c in text if c.isalnum() and not unicodedata.combining(c))


def clean_name(name):
    return ' '.join(str(name).strip().split())


def load_aliases():
    raw = json.loads((ROOT / 'config/player_aliases.json').read_text(encoding='utf-8'))
    return {name_key(k): v for k, v in raw.items()}


def resolve_name(name, aliases):
    cleaned = clean_name(name)
    canonical = aliases.get(name_key(cleaned), cleaned)
    method = 'curated_workbook_variant' if canonical != cleaned else 'normalized_name'
    return name_key(canonical), canonical, method


def register_player(db, name, aliases, preferred=None):
    key, canonical, method = resolve_name(name, aliases)
    display = (preferred or {}).get(key, canonical)
    db.execute('INSERT OR IGNORE INTO players VALUES (?,?,NULL,?)',
               (key, display, 'local_name_identity; NBA ID not yet matched'))
    db.execute('INSERT OR IGNORE INTO player_aliases VALUES (?,?,?)', (name, key, method))
    return key


def season_name(start):
    return f'{start}-{str(start + 1)[-2:]}'


def rows(db, sql, parameters=()):
    cursor = db.execute(sql, parameters) if parameters else db.execute(sql)
    return [{k: float(v) if isinstance(v, Decimal) else v for k,v in dict(r).items()}
            for r in cursor]


def build_catalog(workbook_path, db):
    workbook_path = Path(workbook_path)
    digest = hashlib.sha256(workbook_path.read_bytes()).hexdigest()
    source_id = digest
    if getattr(db, 'dialect', None) == 'postgres':
        db.execute("SELECT pg_advisory_xact_lock(hashtext('survivor-workbook-import'))")
    existing = db.execute('SELECT source_id FROM source_workbooks').fetchall()
    if existing:
        if len(existing) == 1 and existing[0][0] == source_id:
            return source_id
        raise ValueError('A different workbook is already loaded. Revised-workbook reconciliation is required before replacing league history.')
    rules = json.loads((ROOT / 'config/league.json').read_text(encoding='utf-8'))
    aliases = load_aliases()
    wb = openpyxl.load_workbook(workbook_path, data_only=False)
    cached = openpyxl.load_workbook(workbook_path, data_only=True)
    preferred_counts = defaultdict(Counter)
    for sheet in wb:
        tracking = 'tracking' in sheet.title.lower()
        for row in sheet:
            for offset in ([0] if tracking else [0, 3]):
                name = row[offset].value
                if not isinstance(name, str) or not name.strip():
                    continue
                if tracking and row[0].row == 1:
                    continue
                if not tracking and not isinstance(row[offset + 1].value, (int, float)):
                    continue
                key, canonical, _ = resolve_name(name, aliases)
                preferred_counts[key][canonical] += 1
    preferred = {k: c.most_common(1)[0][0] for k, c in preferred_counts.items()}

    def issue(code, detail, season=None, sheet=None, cell=None, severity='review'):
        db.execute('INSERT INTO data_issues(source_id,severity,code,season,sheet,source_cell,detail) VALUES (?,?,?,?,?,?,?)',
                   (source_id, severity, code, season, sheet, cell, detail))

    try:
        with db:
            metadata = {'created': str(wb.properties.created), 'modified': str(wb.properties.modified),
                        'league_settings_at_import': rules,
                        'aliases_sha256': hashlib.sha256((ROOT / 'config/player_aliases.json').read_bytes()).hexdigest()}
            db.execute('INSERT INTO source_workbooks VALUES (?,?,?,?,?)',
                       (source_id, workbook_path.name, digest, now(), json.dumps(metadata)))
            tracking_starts = []
            for sheet in wb:
                if 'tracking' in sheet.title.lower():
                    match = re.search(r'(\d{2})-(\d{2})', sheet.title)
                    if not match:
                        raise ValueError(f'Unknown tracking season: {sheet.title}')
                    start = 2000 + int(match[1])
                    tracking_starts.append(start)
                    db.execute('INSERT INTO seasons VALUES (?,?,?,?)',
                               (season_name(start), start-rules['first_season_start']+1, start, rules['season_mapping_status']))
            if sorted(tracking_starts) != list(range(min(tracking_starts), max(tracking_starts)+1)):
                raise ValueError('Tracking seasons are not consecutive; season mapping needs review.')

            for sheet_index, sheet in enumerate(wb):
                tracking = 'tracking' in sheet.title.lower()
                cells = [c for row in sheet for c in row if c.value is not None or c.comment]
                db.execute('INSERT INTO source_sheets VALUES (?,?,?,?,?,?,?,?)',
                           (source_id, sheet.title, sheet_index, 'tracking' if tracking else 'franchise_history',
                            sheet.max_row, sheet.max_column, len(cells), sheet.sheet_state))
                db.executemany('INSERT INTO raw_cells VALUES (?,?,?,?,?,?,?,?,?,?,?)',
                               [(source_id, sheet.title, c.coordinate, c.row, c.column,
                                 json.dumps(c.value, default=str), json.dumps(cached[sheet.title][c.coordinate].value, default=str),
                                 c.data_type, c.style_id, c.number_format, c.comment.text if c.comment else None) for c in cells])
                if tracking:
                    start = 2000 + int(re.search(r'(\d{2})-(\d{2})', sheet.title)[1])
                    season = season_name(start)
                    for row in sheet.iter_rows(min_row=2):
                        if row[0].value is None:
                            continue
                        name = str(row[0].value)
                        key = register_player(db, name, aliases, preferred)
                        for cell in row[1:]:
                            if cell.value is None or cell.value == '':
                                continue
                            participant = sheet.cell(1, cell.column).value
                            if participant is None:
                                issue('tracking_without_participant', str(cell.value), season, sheet.title, cell.coordinate)
                                continue
                            db.execute('INSERT INTO tracking_entries(source_id,season,player_id,raw_name,participant_raw,marker_raw,marker_number,sheet,source_cell) VALUES (?,?,?,?,?,?,?,?,?)',
                                       (source_id, season, key, name, str(participant), str(cell.value),
                                        cell.value if isinstance(cell.value, (int, float)) else None, sheet.title, cell.coordinate))
                    continue

                team_id = season = league_year = None
                orders = {'opening': 0, 'final': 0}
                for row in sheet:
                    first = row[0].value
                    header = HEADER.match(str(first or ''))
                    if header:
                        year_text = header[1].lower()
                        league_year = {'one': 1, 'two': 2}.get(year_text)
                        if league_year is None:
                            league_year = int(year_text)
                        season = season_name(rules['first_season_start'] + league_year - 1)
                        team_id = f'{season}:{sheet.title}'
                        db.execute('INSERT INTO team_seasons VALUES (?,?,?,?,?,?,?)',
                                   (team_id, source_id, season, sheet.title, int(header[2]), str(first), row[0].row))
                        orders = {'opening': 0, 'final': 0}
                        continue
                    for stage, offset in [('opening', 0), ('final', 3)]:
                        name_cell, cost_cell, year_cell = row[offset:offset+3]
                        name = name_cell.value
                        if name is None or name == '':
                            if cost_cell.value is not None or year_cell.value is not None:
                                issue('orphan_cost_or_year', 'Cost/year without player; preserved in raw_cells.', season, sheet.title, cost_cell.coordinate)
                            continue
                        if str(name).strip().lower() in COLUMN_HEADERS:
                            continue
                        if str(name).strip().lower() == 'vacant':
                            issue('explicit_vacancy', 'Workbook records a vacant final roster slot.', season, sheet.title, name_cell.coordinate, 'info')
                            continue
                        if team_id is None:
                            issue('unparsed_cell', str(name), sheet=sheet.title, cell=name_cell.coordinate)
                            continue
                        orders[stage] += 1
                        raw_cost = cached[sheet.title][cost_cell.coordinate].value
                        cost = raw_cost if isinstance(raw_cost, (int, float)) and raw_cost >= 0 else None
                        raw_year = cached[sheet.title][year_cell.coordinate].value
                        contract = int(raw_year) if isinstance(raw_year, (int, float)) and raw_year >= 0 and int(raw_year) == raw_year else None
                        if cost is None or contract is None:
                            issue('invalid_cost_or_contract', f'cost={raw_cost!r}, year={raw_year!r}', season, sheet.title, name_cell.coordinate)
                        keeper = stage == 'opening' and league_year > 1 and orders[stage] <= rules['keepers_per_team']
                        acquisition = 'keeper' if keeper else ('auction' if stage == 'opening' else 'final_snapshot')
                        basis = 'first two opening players per user; founding season has no keepers' if stage == 'opening' else 'right-hand roster per user; not another sale'
                        if stage == 'opening' and (contract is None or (keeper and contract < 2) or (not keeper and contract != 1)):
                            issue('keeper_position_contract_conflict', f'{acquisition} by position but recorded contract year={contract}', season, sheet.title, name_cell.coordinate)
                        if stage == 'opening' and cost == 0:
                            issue('zero_opening_cost', 'Zero opening cost excluded from clean auction-sales view.', season, sheet.title, cost_cell.coordinate)
                        key = register_player(db, str(name), aliases, preferred)
                        if key == 'jalenjacksonjr':
                            issue('ambiguous_player_name', 'Possibly Jaren Jackson Jr.; retained as a separate identity pending review.', season, sheet.title, name_cell.coordinate)
                        note = row[6].value if stage == 'final' and len(row) > 6 else None
                        db.execute('INSERT INTO roster_entries(team_season_id,player_id,stage,roster_order,raw_name,recorded_cost,contract_year_recorded,acquisition_class,classification_basis,source_cell,cost_cell,contract_year_cell,note_raw) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)',
                                   (team_id, key, stage, orders[stage], str(name), cost, contract, acquisition,
                                    basis, name_cell.coordinate, cost_cell.coordinate, year_cell.coordinate, None if note is None else str(note)))

            for summary in rows(db, 'SELECT * FROM team_summary'):
                if summary['opening_players'] != rules['expected_opening_roster_size']:
                    issue('opening_roster_size', f"Recorded {summary['opening_players']} players; expected {rules['expected_opening_roster_size']} from workbook pattern. Do not invent a missing purchase.", summary['season'], summary['franchise_sheet'])
            for duplicate in rows(db, "SELECT season,player_id,stage,COUNT(*) n FROM roster_history GROUP BY season,player_id,stage HAVING (stage='opening' AND COUNT(*)>1)"):
                locations = rows(db, "SELECT franchise_sheet,source_cell FROM roster_history WHERE season=? AND player_id=? AND stage='opening'", (duplicate['season'], duplicate['player_id']))
                issue('duplicate_opening_player', f"{duplicate['player_id']} appears {duplicate['n']} times: {locations}. Excluded from clean auction_sales.", duplicate['season'])
            for duplicate in rows(db, "SELECT season,franchise_sheet,player_id,stage,COUNT(*) n FROM roster_history GROUP BY season,franchise_sheet,player_id,stage HAVING COUNT(*)>1"):
                issue('duplicate_on_same_roster', f"{duplicate['player_id']} occurs {duplicate['n']} times on the {duplicate['stage']} roster.", duplicate['season'], duplicate['franchise_sheet'])
            for season_row in rows(db, 'SELECT season FROM seasons'):
                season = season_row['season']
                finishes = [r[0] for r in db.execute('SELECT finish FROM team_seasons WHERE season=?', (season,))]
                if sorted(finishes) != list(range(1, rules['teams'] + 1)):
                    issue('finish_order_inconsistent', f'Recorded finishes: {sorted(finishes)}; expected 1 through 15.', season)
            for mismatch in rows(db, "SELECT f.season,f.franchise_sheet,f.player,f.source_cell,f.recorded_cost,a.recorded_cost opening_cost FROM roster_history f JOIN roster_history a ON f.source_id=a.source_id AND f.season=a.season AND f.player_id=a.player_id AND a.stage='opening' WHERE f.stage='final' AND f.recorded_cost<>a.recorded_cost"):
                issue('final_cost_differs_from_opening', f"{mismatch['player']}: final recorded cost {mismatch['recorded_cost']}, opening cost {mismatch['opening_cost']}. Both preserved; final cost is not auction training data.", mismatch['season'], mismatch['franchise_sheet'], mismatch['source_cell'])
            if getattr(db, 'dialect', None) != 'postgres' and db.execute('PRAGMA foreign_key_check').fetchall():
                raise ValueError('Foreign-key validation failed.')
    finally:
        wb.close()
        cached.close()
    return source_id


STAT_HEADERS = ['player','nba_team','positions','games','minutes_pg','pts_pg','reb_pg','ast_pg','stl_pg','blk_pg','fg3m_pg','fgm_pg','fga_pg','ftm_pg','fta_pg']


def export_catalog(db, output_dir):
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    csv_dir = output_dir / 'csv'
    csv_dir.mkdir(exist_ok=True)
    roster_sql = 'SELECT * FROM roster_history'
    season_sql = """SELECT s.season,s.league_year,COUNT(t.team_season_id) team_seasons,
        (SELECT COUNT(*) FROM roster_history h WHERE h.season=s.season AND stage='opening') opening_entries,
        (SELECT COUNT(*) FROM auction_sales a WHERE a.season=s.season) clean_auction_sales,
        (SELECT COUNT(*) FROM keeper_costs k WHERE k.season=s.season) keeper_entries,
        (SELECT COUNT(*) FROM roster_history h WHERE h.season=s.season AND stage='final') final_entries,
        (SELECT SUM(recorded_cost) FROM roster_history h WHERE h.season=s.season AND stage='opening') recorded_opening_spend,
        (SELECT MAX(recorded_cost) FROM auction_sales a WHERE a.season=s.season) highest_clean_auction_price
        FROM seasons s JOIN team_seasons t USING(season) GROUP BY s.season ORDER BY s.league_year"""
    queries = {
        'Seasons': season_sql,
        'Teams': 'SELECT * FROM team_summary ORDER BY season DESC,finish',
        'Opening rosters': roster_sql + " WHERE stage='opening' ORDER BY season DESC,franchise_sheet,roster_order",
        'Auction purchases': 'SELECT * FROM auction_sales ORDER BY season DESC,recorded_cost DESC',
        'Keepers': 'SELECT * FROM keeper_costs ORDER BY season DESC,franchise_sheet,roster_order',
        'Final rosters': roster_sql + " WHERE stage='final' ORDER BY season DESC,franchise_sheet,roster_order",
        'Tracking': 'SELECT t.season,p.display_name AS player,t.raw_name,t.participant_raw,t.marker_raw,t.marker_number,t.sheet,t.source_cell FROM tracking_entries t JOIN players p USING(player_id) ORDER BY t.season DESC,p.display_name,t.marker_number',
        'Players': 'SELECT * FROM players ORDER BY display_name',
        'Name aliases': 'SELECT a.raw_name,p.display_name,a.player_id,a.resolution_method FROM player_aliases a JOIN players p USING(player_id) ORDER BY p.display_name,a.raw_name',
        'Review': 'SELECT severity,code,season,sheet,source_cell,detail FROM data_issues ORDER BY season DESC,code,sheet',
        'Source sheets': 'SELECT sheet,sheet_index,kind,max_row,max_column,populated_cells,hidden_state FROM source_sheets ORDER BY sheet_index',
    }
    summaries = rows(db, season_sql)
    metrics = {
        'source_sheets': db.execute('SELECT COUNT(*) FROM source_sheets').fetchone()[0],
        'seasons': len(summaries),
        'team_seasons': db.execute('SELECT COUNT(*) FROM team_seasons').fetchone()[0],
        'players': db.execute('SELECT COUNT(*) FROM players').fetchone()[0],
        'raw_name_variants': db.execute('SELECT COUNT(*) FROM player_aliases').fetchone()[0],
        'opening_entries': sum(r['opening_entries'] for r in summaries),
        'clean_auction_sales': sum(r['clean_auction_sales'] for r in summaries),
        'keeper_entries': sum(r['keeper_entries'] for r in summaries),
        'final_entries': sum(r['final_entries'] for r in summaries),
        'tracking_entries': db.execute('SELECT COUNT(*) FROM tracking_entries').fetchone()[0],
        'raw_cells': db.execute('SELECT COUNT(*) FROM raw_cells').fetchone()[0],
        'review_issues': db.execute("SELECT COUNT(*) FROM data_issues WHERE severity='review'").fetchone()[0],
        'statistical_datasets_loaded': db.execute('SELECT COUNT(*) FROM stat_datasets').fetchone()[0],
        'valuation_runs': db.execute('SELECT COUNT(*) FROM valuation_runs').fetchone()[0],
    }
    overview = [
        ('Purpose', 'League-history catalog; statistics and projected prices have not yet been loaded or calculated.'),
        ('Database destination', 'Railway Postgres, schema survivor. Preview mode uses memory only; no local database is maintained.'),
        ('Season mapping', 'Year 1 = 2015-16 through Year 11 = 2025-26, inferred from tracking tabs.'),
        ('Keepers', 'First two players on left in Years 2-11; founding year has no keepers. Contract-year column cross-checked.'),
        ('Auction purchases', 'Opening nonkeepers with positive cost, recorded contract year 1, and unique opening player-season.'),
        ('Final rosters', 'Right-hand snapshots at elimination/win, not new auction transactions.'),
        ('Manager history', 'Franchise sheet labels are current labels; do not attribute older spending to a current manager without reconciliation.'),
        ('Tracking markers', 'Preserved exactly; these are not prices or assumed keeper ages.'),
        ('Name matching', 'Conservative text normalization and explicit editable spelling aliases; original names retained; no NBA IDs yet.'),
        ('Unconfirmed league settings', 'Apparent $200 budget and 15-player opening rosters; survivor timing, resets, lineup rules, keeper escalation unknown.'),
        ('Source preservation', 'Original workbook untouched; populated cells and comments archived in raw_cells with source hash.'),
    ] + list(metrics.items())
    out = openpyxl.Workbook()
    ws = out.active
    ws.title = 'Read me'
    ws.append(['Item', 'Value'])
    for pair in overview:
        ws.append(pair)
    for title, query in queries.items():
        cursor = db.execute(query)
        headers = [d[0] for d in cursor.description]
        values = [[r[h] for h in headers] for r in cursor]
        sheet = out.create_sheet(title)
        sheet.append(headers)
        for value in values:
            sheet.append(value)
        with (csv_dir / (title.lower().replace(' ', '_') + '.csv')).open('w', encoding='utf-8-sig', newline='') as stream:
            writer = csv.writer(stream)
            writer.writerow(headers)
            writer.writerows(values)
    template = out.create_sheet('Stats import template')
    template.append(STAT_HEADERS)
    with (output_dir / 'stats_import_template.csv').open('w', encoding='utf-8', newline='') as stream:
        csv.writer(stream).writerow(STAT_HEADERS)
    for sheet in out:
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for c in sheet[1]:
            c.font = Font(color='FFFFFF', bold=True)
            c.fill = PatternFill('solid', fgColor='17365D')
        for index, column in enumerate(sheet.columns, start=1):
            max_len = max(len(str(c.value or '')) for c in column)
            sheet.column_dimensions[get_column_letter(index)].width = min(85, max(14, max_len+2))
        # Treat all workbook strings as text, including any source formula or note.
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
    out.save(output_dir / 'survivor_catalog.xlsx')
    report = {'generated_at': now(), 'metrics': metrics, 'seasons': summaries,
              'issues_by_code': rows(db, 'SELECT code,COUNT(*) AS count FROM data_issues GROUP BY code ORDER BY count DESC')}
    (output_dir / 'catalog_summary.json').write_text(json.dumps(report, indent=2), encoding='utf-8')
    report_lines = ['SURVIVOR HISTORY CATALOG', '', *[f'{k}: {v}' for k,v in metrics.items()], '',
                    'Season       Opening  Auctions  Keepers  Final  Recorded spend']
    report_lines.extend(f"{r['season']:12} {r['opening_entries']:7} {r['clean_auction_sales']:8} {r['keeper_entries']:8} {r['final_entries']:6} {r['recorded_opening_spend']:14}" for r in summaries)
    report_lines += ['', 'Data issues (full details in Review tab and csv/review.csv):']
    report_lines += [f"{r['code']}: {r['count']}" for r in report['issues_by_code']]
    report_lines += ['', *[f'{k}: {v}' for k,v in overview[:11]]]
    (output_dir / 'catalog_summary.txt').write_text('\n'.join(report_lines) + '\n', encoding='utf-8')
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('workbook', nargs='?', default='Survivor keeper log 2025.xlsx')
    parser.add_argument('--preview', action='store_true', help='Validate and export using memory only; do not connect to Postgres.')
    parser.add_argument('--database-url-env', default='DATABASE_URL')
    parser.add_argument('--output', default='output')
    args = parser.parse_args()
    db = preview_database() if args.preview else railway_database(args.database_url_env)
    try:
        build_catalog(args.workbook, db)
        report = export_catalog(db, args.output)
        print(json.dumps(report, indent=2))
    finally:
        db.close()


if __name__ == '__main__':
    main()

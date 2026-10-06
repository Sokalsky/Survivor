"""Fill an Excel workbook from free historical stats and public projections."""
from __future__ import annotations

import argparse
from collections import defaultdict
import csv
from datetime import datetime, timezone
import hashlib
from html.parser import HTMLParser
import json
import math
from pathlib import Path
import re
import time
from urllib.request import Request, urlopen

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from survivor.catalog import ROOT, STAT_HEADERS, build_catalog, load_aliases, resolve_name, rows, season_name
from survivor.database import preview_database, railway_database
from survivor.import_stats import import_stats, validate_csv
from survivor.keepers import keeper_summary, sync_bundled_keepers

PROJECTION_URL = 'https://www.fantasypros.com/nba/projections/overall.php'
MISSING_SHOOTING = 'Missing projected FGM, FGA, FTM and FTA; not ready for eight-category valuation.'
TOTAL_FIELDS = {'minutes_pg':'mp', 'pts_pg':'pts', 'reb_pg':'trb', 'ast_pg':'ast',
                'stl_pg':'stl', 'blk_pg':'blk', 'fg3m_pg':'fg3', 'fgm_pg':'fg',
                'fga_pg':'fga', 'ftm_pg':'ft', 'fta_pg':'fta'}


class TableParser(HTMLParser):
    """Read a specific public HTML table, retaining source identifiers."""
    def __init__(self, table_id):
        super().__init__(convert_charrefs=True)
        self.table_id = table_id
        self.depth = 0
        self.records = []
        self.row = None
        self.cell = None

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag == 'table':
            if self.depth or attrs.get('id') == self.table_id:
                self.depth += 1
        if not self.depth:
            return
        if tag == 'tr':
            self.row = []
        elif tag in ('th', 'td') and self.row is not None:
            self.cell = {'attrs': attrs, 'text': '', 'links': []}
        elif tag == 'a' and self.cell is not None:
            self.cell['links'].append(attrs)

    def handle_data(self, data):
        if self.depth and self.cell is not None:
            self.cell['text'] += data

    def handle_endtag(self, tag):
        if not self.depth:
            return
        if tag in ('td', 'th') and self.cell is not None:
            self.cell['text'] = ' '.join(self.cell['text'].split())
            self.row.append(self.cell)
            self.cell = None
        elif tag == 'tr' and self.row is not None:
            if self.row:
                self.records.append(self.row)
            self.row = None
        elif tag == 'table':
            self.depth -= 1

    def handle_comment(self, data):
        # Sports Reference sometimes wraps entire tables in HTML comments.
        if not self.depth and f'id="{self.table_id}"' in data:
            nested = TableParser(self.table_id)
            nested.feed(data)
            self.records.extend(nested.records)


def numeric(value, label):
    try:
        result = float(value.replace(',', ''))
    except (TypeError, ValueError):
        raise ValueError(f'Missing or invalid {label}: {value!r}') from None
    if not math.isfinite(result) or result < 0:
        raise ValueError(f'Invalid {label}: {value!r}')
    return result


def parse_history(html, season):
    expected_title = f'{season} NBA Player Stats: Totals'
    if expected_title not in html:
        raise ValueError(f'Historical source did not identify the requested {season} totals.')
    table = TableParser('totals_stats')  # Excludes the separate postseason table.
    table.feed(html)
    groups = defaultdict(list)
    for cells in table.records:
        item = {cell['attrs'].get('data-stat'): cell for cell in cells}
        player = item.get('name_display') or item.get('player')
        if not player:
            continue
        source_id = player['attrs'].get('data-append-csv')
        if not source_id:
            continue  # Column headers and league totals.
        get = lambda *keys: next((item[key]['text'] for key in keys if key in item), '')
        games = numeric(get('games', 'g'), 'games')
        if games <= 0 or not games.is_integer():
            raise ValueError(f'Invalid game count for {source_id}.')
        record = {'player': player['text'].rstrip('*'), 'source_player_id': source_id,
                  'nba_team': get('team_name_abbr', 'team_id'), 'positions': get('pos'),
                  'games': int(games), 'age': numeric(get('age'), 'age')}
        for field, source in TOTAL_FIELDS.items():
            record[field.removesuffix('_pg') + '_total'] = numeric(get(source), source)
        groups[source_id].append(record)
    if not groups:
        raise ValueError(f'No regular-season player totals found for {season}.')
    results = []
    for source_id, candidates in groups.items():
        combined = [row for row in candidates if row['nba_team'] == 'TOT' or re.fullmatch(r'\d+TM', row['nba_team'])]
        if len(combined) > 1:
            raise ValueError(f'Multiple aggregate rows for {source_id}.')
        if len(candidates) > 1 and not combined:
            raise ValueError(f'Team splits without a combined row for {source_id}; review required.')
        row = dict(combined[0] if combined else candidates[0])
        row['source_rows'] = len(candidates)
        for field in TOTAL_FIELDS:
            row[field] = row[field.removesuffix('_pg')+'_total'] / row['games']
        if row['fgm_pg'] > row['fga_pg'] or row['ftm_pg'] > row['fta_pg'] or row['fg3m_pg'] > row['fgm_pg']:
            raise ValueError(f'Invalid shooting totals for {source_id}.')
        row['fg_pct'] = row['fgm_pg']/row['fga_pg'] if row['fga_pg'] else None
        row['ft_pct'] = row['ftm_pg']/row['fta_pg'] if row['fta_pg'] else None
        row['season'] = season
        results.append(row)
    return sorted(results, key=lambda row: row['player'])


def parse_projections(html, season):
    if f'Overall {season} Projections' not in html:
        raise ValueError('Projection source season does not match the requested season.')
    published = re.search(r'<time\b[^>]*datetime=["\'](\d{4}-\d{2}-\d{2})', html)
    if not published:
        raise ValueError('Projection publication date is missing.')
    published_date = published.group(1)
    table = TableParser('data')
    table.feed(html)
    expected = ['Player','PTS','REB','AST','BLK','STL','FG%','FT%','3PM','GP','MIN','TO']
    if not table.records or [cell['text'] for cell in table.records[0]] != expected:
        raise ValueError('The public projection table columns changed; review the adapter.')
    results, seen = [], set()
    for cells in table.records[1:]:
        if [cell['text'] for cell in cells] == expected:
            continue
        if len(cells) != len(expected):
            raise ValueError('Incomplete projection row.')
        source = next((link for link in cells[0]['links'] if link.get('fp-player-name')), None)
        if not source:
            raise ValueError('Projection row has no player identity.')
        source_id = source['href'].rsplit('/', 1)[-1].removesuffix('.php')
        if source_id in seen:
            raise ValueError(f'Duplicate projected player: {source_id}')
        seen.add(source_id)
        attributes = re.search(r'\(([^()]+?)\s+-\s+([^()]+)\)', cells[0]['text'])
        if not attributes:
            raise ValueError(f'Missing team/position for {source_id}.')
        values = {key:numeric(cell['text'], key) for key,cell in zip(expected[1:],cells[1:])}
        gp = values['GP']
        if gp <= 0 or gp > 85 or values['FG%'] > 1 or values['FT%'] > 1:
            raise ValueError(f'Invalid projected games or percentage for {source_id}.')
        row = {field:None for field in STAT_HEADERS}
        row.update(player=source['fp-player-name'], source_player_id=source_id,
                   season=season, nba_team=attributes.group(1), positions=attributes.group(2),
                   games=gp, fg_pct=values['FG%'], ft_pct=values['FT%'],
                   quality_note=MISSING_SHOOTING, ready_for_valuation=False,
                   source_date=published_date)
        for field, label in [('pts_pg','PTS'),('reb_pg','REB'),('ast_pg','AST'),('blk_pg','BLK'),
                             ('stl_pg','STL'),('fg3m_pg','3PM'),('minutes_pg','MIN')]:
            row[field] = values[label]/gp
        results.append(row)
    if not results:
        raise ValueError('The public projection table contains no players.')
    return results, published_date


def download(url, cache_dir, key, *, refresh=False):
    html_path = cache_dir / (key+'.html')
    meta_path = cache_dir / (key+'.json')
    if html_path.exists() and meta_path.exists() and not refresh:
        content = html_path.read_bytes()
        metadata = json.loads(meta_path.read_text(encoding='utf-8'))
        if metadata['url'] != url or metadata['sha256'] != hashlib.sha256(content).hexdigest():
            raise ValueError('Cached source does not match its provenance.')
    else:
        request = Request(url, headers={'User-Agent':'SurvivorStats/1.0 (personal fantasy basketball research)'})
        # No credential, subscription, or paid API is used. HTTP errors are surfaced.
        with urlopen(request, timeout=30) as response:
            content = response.read()
        metadata = {'url':url, 'retrieved_at':datetime.now(timezone.utc).isoformat(),
                    'sha256':hashlib.sha256(content).hexdigest()}
        cache_dir.mkdir(parents=True, exist_ok=True)
        html_path.write_bytes(content)
        meta_path.write_text(json.dumps(metadata, indent=2), encoding='utf-8')
        time.sleep(3.2)  # Less than 20 page requests/minute, cached for reruns.
    return content.decode('utf-8'), metadata


def write_csv(path, records, headers):
    with path.open('w', encoding='utf-8-sig', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=headers, extrasaction='ignore')
        writer.writeheader()
        writer.writerows(records)


def add_sheet(workbook, title, records, headers):
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for record in records:
        sheet.append([record.get(key) for key in headers])
    for cell in sheet[1]:
        cell.fill = PatternFill('solid', fgColor='20352B')
        cell.font = Font(color='FFFFFF', bold=True)
        cell.alignment = Alignment(wrap_text=True)
    sheet.row_dimensions[1].height = 30
    for index, name in enumerate(headers, 1):
        width = 29 if name in ('player','raw_name') else 38 if name in ('quality_note','note','url') else 19
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for name, cell in zip(headers, row):
            if isinstance(cell.value, str):
                cell.data_type = 's'  # Source text must never become an Excel formula.
            if name.endswith('_pct'):
                cell.number_format = '0.0%'
            elif name.endswith('_pg'):
                cell.number_format = '0.00'
            elif name in ('keeper_cost','remaining_budget','recorded_cost'):
                cell.number_format = '$0.00'
    sheet.freeze_panes = 'D2' if len(headers)>8 else 'A2'
    sheet.sheet_view.showGridLines = False
    if title == 'Read me':
        sheet.column_dimensions['A'].width = 28
        sheet.column_dimensions['B'].width = 110
        for row in sheet.iter_rows(min_row=2):
            row[1].alignment = Alignment(wrap_text=True, vertical='top')
            sheet.row_dimensions[row[1].row].height = 46
    if records:
        table = Table(displayName='Stats'+str(len(workbook.worksheets)), ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium4', showRowStripes=True)
        sheet.add_table(table)
    return sheet


def build_workbook(output_dir, first_start, last_start, target_season, *, refresh=False, include_projections=True):
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    cache = output_dir/'cache'
    datasets, sources, errors = [], [], []
    aliases = load_aliases()
    for start in range(first_start, last_start+1):
        season = season_name(start)
        url = f'https://www.basketball-reference.com/leagues/NBA_{start+1}_totals.html'
        try:
            html, metadata = download(url, cache, 'actuals-'+season, refresh=refresh)
            records = parse_history(html, season)
            if len(records) < 350:
                raise ValueError(f'Only {len(records)} player rows found; refusing an apparently partial NBA season.')
            path = output_dir / ('actuals-'+season+'.csv')
            write_csv(path, records, STAT_HEADERS)
            validate_csv(path)
            source = dict(kind='actual', season=season, source_name='Basketball Reference',
                          as_of_date=metadata['retrieved_at'][:10], file=path.name,
                          players=len(records), ready_for_import=True, **metadata)
            sources.append(source)
            datasets.append((season, records))
            print(f'{season}: {len(records)} historical players', flush=True)
        except Exception as exc:
            errors.append({'season':season, 'source':'Basketball Reference', 'note':str(exc)})
            print(f'{season}: FAILED - {exc}', flush=True)
    projections = []
    if include_projections:
        try:
            html, metadata = download(PROJECTION_URL, cache, 'projections-'+target_season, refresh=refresh)
            projections, published_date = parse_projections(html, target_season)
            if len(projections) < 225:
                raise ValueError(f'Only {len(projections)} projections found; fewer than the observed 225 league roster spots.')
            source = dict(kind='projection', season=target_season, source_name='FantasyPros public consensus',
                          as_of_date=published_date, file='projections-'+target_season+'-INCOMPLETE.csv',
                          players=len(projections), ready_for_import=False, note=MISSING_SHOOTING, **metadata)
            sources.append(source)
            write_csv(output_dir/source['file'], projections,
                      STAT_HEADERS+['fg_pct','ft_pct','source_date','quality_note'])
            print(f'{target_season}: {len(projections)} projections; shooting-volume fields unavailable', flush=True)
        except Exception as exc:
            projections = []
            errors.append({'season':target_season,'source':'FantasyPros','note':str(exc)})
            print(f'Projections: FAILED - {exc}', flush=True)

    db = preview_database()
    try:
        build_catalog(ROOT/'Survivor keeper log 2025.xlsx', db)
        sync_bundled_keepers(db)
        catalog = {row['player_id']:row['display_name'] for row in rows(db,'SELECT player_id,display_name FROM players')}
        prices = rows(db, 'SELECT season,player,player_id,franchise_sheet,recorded_cost FROM auction_sales ORDER BY season,recorded_cost DESC')
        draft = keeper_summary(db, target_season)
    finally:
        db.close()
    keepers = {row['player_id']:row for row in draft['rows']} if draft else {}
    for dataset_label, records in datasets+[('projection',projections)]:
        for record in records:
            key, _, _ = resolve_name(record['player'], aliases)
            record['local_player_id'] = key
            record['archive_match'] = 'matched by normalized name' if key in catalog else 'not in price archive / review name'
            if dataset_label == 'projection':
                keeper = keepers.get(key)
                record['draft_status'] = 'kept' if keeper else 'available' if draft else 'unknown'
                record['keeper_franchise'] = keeper['franchise'] if keeper else None
                record['keeper_cost'] = keeper['keeper_cost'] if keeper else None

    workbook = Workbook()
    workbook.remove(workbook.active)
    notes = [
        {'item':'Purpose','note':'Free source data for Survivor comparables. This workbook does not calculate projected dollar values.'},
        {'item':'History','note':f'{season_name(first_start)} through {season_name(last_start)} requested; regular season only, one combined row per player-season.'},
        {'item':'Historical units','note':'Per-game values calculated from exact source totals. Aggregate TOT / multi-team rows replace individual team splits.'},
        {'item':'Historical dates','note':'As-of dates are retrieval dates, not evidence that final stats were known before that season began.'},
        {'item':'Projection units','note':'FantasyPros season totals divided by projected games; published FG% and FT% retained separately.'},
        {'item':'Projection completeness','note':MISSING_SHOOTING+' Empty cells are not zeroes. No shooting volume has been invented.'},
        {'item':'Identity matching','note':'Source player IDs and normalized archive matches are included. Unmatched names require review before modeling.'},
        {'item':'Keepers','note':'Confirmed keeper availability applies only to the target projection season. Historical purchases remain separate.'},
        {'item':'Refresh','note':'Run python -m survivor.stats_workbook. Use --refresh to download again; cache snapshots retain source hashes and dates.'},
        {'item':'Database','note':'Validated actuals CSVs can be imported with --import-history. Incomplete projections are blocked from that import.'},
        {'item':'Fetch errors','note':str(len(errors))+' failed sources; see Issues and manifest.json. A failed refresh is never labeled complete.'},
    ]
    add_sheet(workbook,'Read me',notes,['item','note'])
    history_headers = ['season','player','source_player_id','local_player_id','archive_match','nba_team','positions','age']+STAT_HEADERS[3:]+['fg_pct','ft_pct']+[field.removesuffix('_pg')+'_total' for field in TOTAL_FIELDS]+['source_rows']
    for season, records in datasets:
        add_sheet(workbook, 'Stats '+season, records, history_headers)
    if include_projections:
        add_sheet(workbook, 'Projections '+target_season, projections,
                  ['player','source_player_id','local_player_id','archive_match','nba_team','positions']+STAT_HEADERS[3:]+['fg_pct','ft_pct','source_date','draft_status','keeper_franchise','keeper_cost','ready_for_valuation','quality_note'])
    add_sheet(workbook,'Auction prices',prices,['season','player','player_id','franchise_sheet','recorded_cost'])
    covered = {season:{row['local_player_id'] for row in records} for season,records in datasets}
    coverage = []
    for price in prices:
        prior = season_name(int(price['season'][:4])-1)
        coverage.append({**price,
            'same_season_stats':'matched' if price['player_id'] in covered.get(price['season'],set()) else 'no match / no games / source missing',
            'prior_season_stats':'matched' if price['player_id'] in covered.get(prior,set()) else 'no match / rookie / no games / source missing'})
    add_sheet(workbook,'Price stat coverage',coverage,['season','player','player_id','franchise_sheet','recorded_cost','same_season_stats','prior_season_stats'])
    if draft:
        add_sheet(workbook,'Keepers '+target_season,draft['rows'],['franchise','player','player_id','keeper_cost'])
        add_sheet(workbook,'Team budgets',draft['teams'],['franchise','budget_per_team','keeper_spend','remaining_budget'])
    add_sheet(workbook,'Sources',sources,['kind','season','source_name','as_of_date','retrieved_at','players','ready_for_import','file','url','sha256','note'])
    add_sheet(workbook,'Issues',errors,['season','source','note'])
    workbook_path = output_dir/'survivor-player-stats.xlsx'
    temporary = workbook_path.with_suffix('.tmp.xlsx')
    workbook.save(temporary)
    temporary.replace(workbook_path)
    manifest = {'created_at':datetime.now(timezone.utc).isoformat(), 'workbook':workbook_path.name,
                'historical_seasons':len(datasets), 'historical_rows':sum(len(r) for _,r in datasets),
                'projection_rows':len(projections), 'projections_ready_for_valuation':False,
                'sources':sources,'errors':errors}
    (output_dir/'manifest.json').write_text(json.dumps(manifest,indent=2),encoding='utf-8')
    print(f'Workbook: {workbook_path.resolve()}', flush=True)
    return manifest


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output-dir',type=Path,default=ROOT/'output/stats')
    parser.add_argument('--first-start',type=int,default=2014)
    parser.add_argument('--last-start',type=int,default=2025)
    parser.add_argument('--target-season',default='2026-27')
    parser.add_argument('--refresh',action='store_true')
    parser.add_argument('--history-only',action='store_true')
    parser.add_argument('--import-history',action='store_true',help='Also import complete historical CSVs into Railway Postgres.')
    args = parser.parse_args()
    if not 2000 <= args.first_start <= args.last_start < 2100:
        parser.error('Invalid historical season range.')
    if not re.fullmatch(r'20\d{2}-\d{2}',args.target_season) or args.target_season != season_name(int(args.target_season[:4])):
        parser.error('Target season must be YYYY-YY.')
    manifest = build_workbook(args.output_dir,args.first_start,args.last_start,args.target_season,
                              refresh=args.refresh,include_projections=not args.history_only)
    if args.import_history:
        if manifest['errors']:
            raise SystemExit('Import skipped: resolve source errors before importing this batch.')
        db = railway_database()
        try:
            for source in manifest['sources']:
                if source['kind'] == 'actual' and source['ready_for_import']:
                    import_stats(db,args.output_dir/source['file'],kind='actual',season=source['season'],
                                 source_name=source['source_name'],as_of_date=source['as_of_date'],
                                 source_url=source['url'],notes='Regular-season source totals, converted to per-game values. Retrieval date is not a preseason snapshot.')
        finally:
            db.close()
    if manifest['errors']:
        raise SystemExit(1)


if __name__ == '__main__':
    main()
